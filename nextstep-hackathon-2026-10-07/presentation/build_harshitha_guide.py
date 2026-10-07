from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = Path(__file__).parent / 'HARSHITHA-DEMO-GUIDE.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.7)
sec.left_margin = sec.right_margin = Inches(.8)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    st = doc.styles[name]
    st.font.name = 'Calibri'
    st.font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(27)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(13)
for name in ['Heading 1', 'Heading 2']:
    doc.styles[name].paragraph_format.space_before = Pt(11)
    doc.styles[name].paragraph_format.space_after = Pt(5)
doc.core_properties.title = 'NextStep demo and presentation guide'
doc.core_properties.author = 'NextStep team'

def p(text, style=None):
    doc.add_paragraph(text, style)
def h(text):
    doc.add_heading(text, 2)
def page(title):
    doc.add_page_break()
    doc.add_heading(title, 1)
def cue(time, action, speech):
    h(time)
    x = doc.add_paragraph()
    x.add_run('On screen: ').bold = True
    x.add_run(action)
    p('Harshitha: “' + speech + '”')

doc.add_heading('NextStep demo and presentation guide', 0)
p('Harshitha • BuildToShip • Personalized AI Experiences', 'Subtitle')
p('Present one student journey: start a manageable DSA mission, get help when it feels difficult, and rebuild the remaining plan when the week changes. Use this guide to rehearse a four-minute presentation and record the required three-to-five-minute demo.')
h('The message to remember')
p('“NextStep helps students find a doable next action and return to preparation after an interruption.”')
h('Our real starting point')
p('Darshit described researching GSoC on day one and then not continuing. He also described finding a daily tool habit easier to sustain. Use the first experience as the origin story, with his agreement. It is a personal anecdote, not evidence about all students or proof that NextStep changes behavior.')
p('Suggested opening: “Our teammate Darshit wanted to prepare for GSoC. He researched where to start, but the preparation stopped there. That made us ask: what would make the next action easier to begin?”')
h('The student we demonstrate')
p('Meet Ananya, an illustrative B.Tech CSE student preparing for placements. She plans 60 minutes of DSA on Monday, Wednesday and Friday. Then lab submissions reduce those slots to 30 minutes. She also finds the next mission difficult. This is a realistic scenario, not an interviewed user or testimonial.')
h('What the MVP actually covers')
p('One curated DSA starter track: 12 missions of 30 minutes each. Goal and availability setup, self-reported completion, AI guidance within the next mission, a recovery preview, explicit acceptance and saved progress. Six hours is the catalogue workload estimate, not a promise of placement readiness.')
p('Gemini adapts the explanation and practice steps. Backend rules calculate the schedule. Supabase stores the accepted state. GSoC inspired the story; a GSoC track is not part of this MVP.')
h('Before using the script')
p('At preparation time, the workspace contains an early frontend scaffold; the reviewed handoffs do not establish a working backend or deployment. The following is a rehearsal script for the intended flow. Darshit and Sankirth must verify each feature before Harshitha describes it as working. Use the readiness check on page 5.')

page('Four minute script first half')
p('Harshitha narrates. Darshit operates the app. Sankirth stays ready for integration questions. These time windows include clicks and short pauses; rehearse against the actual app.')
cue('0:00 to 0:30  The problem', 'One simple opening slide; no feature list.',
    'Our teammate Darshit wanted to prepare for GSoC. He researched where to start, but the preparation stopped there. The goal was clear; the next action still felt too big. We built NextStep around that gap: helping a student begin a manageable task and return when their week changes.')
cue('0:30 to 0:55  The student', 'Switch to the app, already signed in with synthetic demo data.',
    'For our demo, imagine Ananya, a CSE student starting placement preparation. She can study for an hour on Monday, Wednesday and Friday. NextStep uses that availability to schedule a small DSA starter track. This example is fictional, but the college workload it represents will feel familiar.')
cue('0:55 to 1:25  A manageable next action', 'Show availability briefly, then the mission title, steps and completion criteria. Use the actual mission text.',
    'Instead of asking her to finish all of DSA, the dashboard gives her one 30-minute mission, with a purpose, concrete steps and a way to check her work. We have curated twelve starter missions. This is a starting sequence, not a claim that six hours makes someone interview-ready.')
cue('1:25 to 1:45  Save progress', 'On a prepared account, record one mission as complete and refresh after success.',
    'We will simulate Ananya marking this practice session complete. This is self-reported progress, not automatic proof that she solved a problem. After refreshing, the completed mission remains saved. That gives the next interaction the context of what she has already done.')
cue('1:45 to 2:15  Personalized AI guidance', 'On the next eligible mission choose too_difficult; enter a short, relevant difficulty. Generate guidance.',
    'Now Ananya is stuck on the next mission. She describes what confused her. Our backend sends the relevant mission and feedback to Gemini. The response suggests a different way to approach the same session, with smaller steps and a check question. The student still decides what to do and when to mark it complete.')
h('Operating notes')
p('Do not invent exact AI wording in advance. Point to one actual step and explain how it addresses the entered difficulty. If the mission is about arrays, a suitable input is “I can read the example, but I do not understand how the index changes.” Use different wording if the selected mission covers another topic.')
p('Do not pretend to complete 30 minutes of study during the presentation. Call the completion action a simulation. If the AI is still loading, pause briefly rather than clicking repeatedly.')

page('Four minute script second half')
cue('2:15 to 2:55  The week changes', 'Open recovery. Change the three 60-minute days to 30 minutes. Preview; point to actual finish date and preserved completions.',
    'Lab submissions now cut Ananya’s available time in half. NextStep shows a revised schedule before changing anything. The remaining workload stays visible, and completed work stays completed. In this flexible-deadline example, the student can see whether the finish date needs to move. She accepts the plan only after seeing the consequence.')
cue('2:55 to 3:15  Confirm the recovery', 'Apply once, wait for success, refresh and show the resulting plan.',
    'The accepted schedule is saved. After a refresh, we still have the completed mission and the revised remaining plan. Recovery does not mean restarting the course or pretending missed work disappeared.')
cue('3:15 to 3:40  How it works', 'One architecture slide; keep infrastructure dashboards closed.',
    'We implement with Antigravity. Vercel serves the interface, Render runs our API, scheduling rules and Gemini requests, and Supabase stores each user’s goal and progress. AI adapts the guidance; ordinary code handles dates and completion records. That separation makes the core scheduling behavior testable.')
cue('3:40 to 4:00  Close', 'Return to the dashboard. Show a tested app URL or QR code if available.',
    'Our focus is the moment a student would otherwise stop: a difficult mission or a disrupted week. NextStep gives them a practical next action and a way back. Our next step is a small student pilot to test whether this helps people return to practice. Thank you.')
h('Five simple slides')
for text in [
    '1. The problem: one sentence from Darshit’s experience. Avoid invented dropout statistics.',
    '2. Meet the demo student: Ananya’s goal and weekly availability. Label it an example.',
    '3. Product demonstration: switch to the actual app. Do not narrate over a wall of screenshots.',
    '4. How it works: Vercel → Render → Supabase and Gemini. Show Antigravity separately as the development tool.',
    '5. What we proved and what comes next: list only verified functionality, then the proposed student pilot.'
]: p(text)
h('If you have only one minute')
p('“Students can know their goal and still struggle to begin or restart. NextStep turns one DSA starter track into manageable missions using their available time. Here is the next mission. When a student struggles, Gemini changes the guidance within that session. When availability changes, the app previews a revised schedule without erasing completed work. The student accepts it, and the saved state survives a refresh. We are testing a focused idea: making the return to preparation easier.” Use only demonstrated claims.')

page('Judge questions and honest answers')
qa = [
('Where is the AI?', 'Gemini receives the next mission and relevant student feedback through Render. It produces structured guidance, an explanation and a check question. It does not calculate authoritative schedules or mark work complete.'),
('Why not ask a chatbot for a roadmap?', 'A chat response can suggest a plan. Our MVP maintains a saved goal, available-time constraints and completion state, then lets the student preview and accept recovery. The useful demonstration is that connected workflow.'),
('What makes this different from a task manager?', 'Our focus is a curated preparation sequence with prerequisites, contextual help and recovery after changed availability. We are not claiming these individual features are new; we need student testing to learn whether the combined experience helps.'),
('How does this fit Personalized AI Experiences?', 'The guidance uses a student’s mission and difficulty feedback rather than a generic answer. The rest of the experience also responds to saved progress and changing availability. Show the actual input and response.'),
('How do you know students learned anything?', 'We do not infer mastery from a completion click. Progress is self-reported practice. Assessing learning would require a separate evaluation that is outside this MVP.'),
('What if the deadline is impossible?', 'The scheduler keeps the workload and shows overflow. A fixed deadline remains fixed. A flexible deadline can move only after the student previews and accepts the revised plan.'),
('What happens when AI fails?', 'We retain the saved goal and show a retryable error. Standard mission content is still available. We do not label hardcoded advice as a successful AI response.'),
('Why only twelve missions?', 'We chose one small sequence to demonstrate the whole experience within the hackathon. Broader DSA coverage and GSoC tracks need additional content review.'),
('Is this proven to improve consistency?', 'No long-term outcome has been measured yet. We propose observing whether consenting students begin a mission, return after an interruption and continue practice. We will report results only after that pilot.'),
('How is student data protected?', 'Our intended design uses authenticated requests and Supabase row ownership policies. Sankirth can explain which isolation checks were actually run. We keep AI credentials on the backend and send only relevant task context to the model.')
]
for question, answer in qa:
    h(question)
    p(answer)

page('Rehearsal and demo readiness')
h('One short team rehearsal')
p('Harshitha leads narration and keeps time. Darshit performs the exact clicks. Sankirth verifies integrations beforehand and takes technical questions when needed. Agree on one quiet cue for switching from slides to app. Avoid all three people narrating the same action.')
h('Developer readiness check')
for text in [
    'Confirm the exact release and mode: fixture, local live or deployed live. Use synthetic accounts and data.',
    'Verify sign-in, goal load, completion save and refresh. Prepare a separate demo account so earlier rehearsals do not exhaust the one-goal flow.',
    'Verify a real Gemini request and a response that addresses the selected mission. Confirm that the API key and model are available.',
    'Verify recovery preview, Apply and refresh. Note the expected actual dates; do not promise a finish date from this script.',
    'Verify the public app URL and direct-route refresh. Open the backend shortly before the demo and check it is responsive.',
    'Give Harshitha a written list of working features and limitations. A green planning-kit check is not an application test.'
]: p('• ' + text)
h('If something fails on stage')
p('AI failure: “The AI service did not return successfully, so the app is keeping the existing plan. I will show the recovery workflow next.” If available, show an earlier genuine AI recording and label it as recorded.')
p('Deployment/network failure: “The live service is unavailable right now. This recording shows the flow we verified earlier.” State whether that recording was local or deployed. If no verified recording exists, show the design as a prototype and say what remains unfinished.')
h('Recording and submission')
p('Record a 3–5 minute video after verification, following the four-minute script. Use readable browser zoom, clear audio and deliberate clicks. Hide passwords, tokens, private tabs and notifications. Keep a local video backup. Check the final video plays from start to finish.')
p('Submit the actual problem statement, solution description with the AI role, public GitHub repository with setup README, working application URL and demo video. Confirm the organizer’s submission method and deadline. Do not display an untested QR code.')
h('Three conversations before judging')
p('If time permits, ask a few consenting classmates: “When did you last stop preparing?”, “What made restarting difficult?” and “What would you expect after pressing Recover?” Note their actual answers and the actual number interviewed. Do not convert a few conversations into a claim of proven demand or measured effectiveness.')
h('Keep these claims accurate')
p('Say “one DSA starter track,” “self-reported progress,” and “AI-assisted guidance.” Do not promise automatic LeetCode verification, a shipped GSoC track, guaranteed placement, measured retention improvement or a complete personalized curriculum.')
p('Basis: the locked NextStep master prompt and official Personalized AI Experiences challenge, linked from Hackathon Themes.xlsx, Sheet1 A14:B14. This guide is a rehearsal aid; developers supply the final evidence of what works.')

footer = sec.footer.paragraphs[0]
footer.alignment = 2
footer.add_run('NextStep  |  ')
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
for run in footer.runs:
    run.font.size = Pt(9)
doc.save(out)
lines = []
for paragraph in doc.paragraphs:
    text = paragraph.text.strip()
    if not text:
        continue
    style = paragraph.style.name
    prefix = '# ' if style == 'Title' else '## ' if style == 'Heading 1' else '### ' if style == 'Heading 2' else ''
    text = text.replace('on page 5', 'under Rehearsal and demo readiness')
    lines.append(prefix + text)
out.with_suffix('.md').write_text('\n\n'.join(lines) + '\n', encoding='utf-8')
print(out)
